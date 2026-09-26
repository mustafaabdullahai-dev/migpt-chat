"""Document ingestion pipeline: parse -> clean -> chunk -> embed -> FAISS + DB."""

from __future__ import annotations

import csv
import hashlib
import io

from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.logging import get_logger
from app.db import crud
from app.rag.embeddings import get_embeddings
from app.rag.vectorstore import get_vectorstore
from app.services.cache import cache_key, get_cache

logger = get_logger("rag.ingestion")


def parse_file(filename: str, content: bytes) -> list[tuple[int, str]]:
    """Returns list of (page_number, page_text) for a supported file type."""
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _parse_pdf(content)
    if ext == "docx":
        return _parse_docx(content)
    if ext in ("csv", "tsv"):
        return _parse_csv(content)
    # txt / md
    text = _decode(content)
    return [(0, text)]


def parse_file_cached(filename: str, content: bytes) -> list[tuple[int, str]]:
    """Identical-document processing cache: the same bytes + filename parse once.

    The parsed pages are stored as JSON-safe [[page, text], …] so the same key
    works for both in-memory and Redis backends.
    """
    cache = get_cache()
    if cache is None:
        return parse_file(filename, content)
    digest = hashlib.sha256(content).hexdigest()
    key = cache_key("parse", filename, digest)
    hit = cache.get(key)
    if hit is not None:
        return _as_pages(hit)
    pages = parse_file(filename, content)
    cache.set(key, [[p, t] for p, t in pages], ttl=3600)
    return pages


def _as_pages(raw: object) -> list[tuple[int, str]]:
    if not isinstance(raw, list):
        return []
    out: list[tuple[int, str]] = []
    for item in raw:
        if isinstance(item, (list, tuple)) and len(item) == 2:
            try:
                out.append((int(item[0]), str(item[1])))
            except (TypeError, ValueError):
                continue
    return out


def ingest_bytes(db: Session, user_id: str, filename: str, content: bytes) -> dict:
    pages = parse_file_cached(filename, content)
    doc = crud.create_document(db, user_id, filename, _content_type(filename), len(content))

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        length_function=len,
    )

    embeddings = get_embeddings()
    vectorstore = get_vectorstore()

    chunk_rows: list[tuple[int, str, dict]] = []
    chunk_ids: list[str] = []
    chunk_texts: list[str] = []

    for page_idx, page_text in pages:
        cleaned = _clean_text(page_text)
        if not cleaned:
            continue
        for chunk in text_splitter.split_text(cleaned):
            cid = crud.new_id()
            chunk_ids.append(f"{doc.document_id}::{cid}")
            chunk_texts.append(chunk)
            chunk_rows.append(
                (
                    page_idx,
                    chunk,
                    {
                        "document_id": doc.document_id,
                        "filename": filename,
                        "source": filename.rsplit(".", 1)[0].lower() if "." in filename else filename,
                        "page": page_idx,
                        "chunk_id": cid,
                    },
                )
            )

    if chunk_texts:
        vectors = embeddings.embed_documents(chunk_texts)
        vectorstore.add(chunk_ids, vectors)
        crud.save_chunks(db, doc.document_id, chunk_rows)
        crud.patch_document_chunk_count(db, doc.document_id, len(chunk_rows))

    logger.info("document ingested", extra={"event": "rag.ingest", "doc": doc.document_id, "chunks": len(chunk_rows)})
    return {"document_id": doc.document_id, "chunks": len(chunk_rows)}


def remove_document_vectors(doc_ids: list[str]) -> None:
    vs = get_vectorstore()
    vs.remove_document(set(doc_ids))


def _clean_text(text: str) -> str:
    import re

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode(content: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace")


def _parse_pdf(content: bytes) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((i, text.strip()))
    return pages


def _parse_docx(content: bytes) -> list[tuple[int, str]]:
    from docx import Document as DocxDocument

    docx_doc = DocxDocument(io.BytesIO(content))
    text = "\n".join(p.text for p in docx_doc.paragraphs if p.text)
    for table in docx_doc.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    return [(0, text)]


def _parse_csv(content: bytes) -> list[tuple[int, str]]:
    text = _decode(content)
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except Exception:
        return [(0, text)]
    blocks: list[str] = []
    for i, row in enumerate(rows):
        if row and any(cell.strip() for cell in row):
            blocks.append((" | ".join(row)))
    return [(i, block) for i, block in enumerate(blocks) if block]


def _content_type(filename: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".csv": "text/csv",
    }.get("." + filename.rsplit(".", 1)[-1].lower(), "application/octet-stream")